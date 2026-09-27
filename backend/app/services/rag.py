"""Permission-aware hybrid RAG for the reference demo.

The important security property is that the ACL filter is applied in the vector
store query itself, before retrieved chunks are handed to the agent/model. Qdrant
supports payload filters and payload indexes for this pattern. See README for the
production hardening notes.
"""
import hashlib
import re
import time
from pathlib import Path
from typing import Any

import numpy as np
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance, VectorParams, PointStruct, Filter, FieldCondition, MatchAny,
)
from sentence_transformers import SentenceTransformer

from ..config import COLLECTION, EMBEDDING_MODEL, QDRANT_URL


class RAGService:
    def __init__(self):
        self.client = self._connect_qdrant()
        self.encoder = SentenceTransformer(EMBEDDING_MODEL)
        self.dim = self.encoder.get_sentence_embedding_dimension()
        self._ensure_collection()

    @staticmethod
    def _connect_qdrant() -> QdrantClient:
        last = None
        for _ in range(20):
            try:
                c = QdrantClient(url=QDRANT_URL)
                c.get_collections()
                return c
            except Exception as exc:
                last = exc
                time.sleep(1)
        raise RuntimeError(f"Qdrant unavailable: {last}")

    def _ensure_collection(self):
        collections = [c.name for c in self.client.get_collections().collections]
        if COLLECTION not in collections:
            self.client.create_collection(
                collection_name=COLLECTION,
                vectors_config=VectorParams(size=self.dim, distance=Distance.COSINE),
            )
        # ACL is a high-value filter field. Indexing it makes filtered search efficient.
        try:
            self.client.create_payload_index(
                collection_name=COLLECTION,
                field_name="allowed_groups",
                field_schema="keyword",
            )
            self.client.create_payload_index(
                collection_name=COLLECTION,
                field_name="department",
                field_schema="keyword",
            )
        except Exception:
            # Index may already exist.
            pass

    @staticmethod
    def _chunk(text: str, size: int = 140, overlap: int = 25):
        # Word-based chunking keeps the demo deterministic and inexpensive.
        words = text.split()
        chunks = []
        i = 0
        while i < len(words):
            chunk = " ".join(words[i:i + size])
            if chunk.strip():
                chunks.append(chunk)
            i += max(1, size - overlap)
        return chunks

    @staticmethod
    def _read_file(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == ".docx":
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"SHEET: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        rows.append(" | ".join(vals))
            return "\n".join(rows)
        raise ValueError(f"Unsupported file type: {suffix}")

    @staticmethod
    def _metadata_for(filename: str, metadata: dict[str, Any] | None = None):
        metadata = metadata or {}
        lower = filename.lower()
        if "procurement" in lower or "purchase" in lower:
            department = "procurement"
        elif "it_" in lower or "service" in lower:
            department = "it"
        elif "hr_" in lower or "leave" in lower:
            department = "hr"
        elif "finance" in lower or "invoice" in lower or "vendor" in lower:
            department = "finance"
        else:
            department = metadata.get("department", "general")

        confidential = "confidential" in lower or metadata.get("sensitivity") == "confidential"
        allowed_groups = metadata.get("allowed_groups")
        if not allowed_groups:
            allowed_groups = [department, "admins"] if confidential else [department, "employees", "admins"]
        return department, allowed_groups, metadata.get("sensitivity", "confidential" if confidential else "internal")

    def ingest_file(self, path: str | Path, metadata: dict[str, Any] | None = None) -> int:
        path = Path(path)
        text = self._read_file(path)
        if not text.strip():
            return 0
        department, allowed_groups, sensitivity = self._metadata_for(path.name, metadata)
        chunks = self._chunk(text)
        vectors = self.encoder.encode(
            [f"passage: {c}" for c in chunks],
            normalize_embeddings=True,
            batch_size=32,
            show_progress_bar=False,
        )
        points = []
        for chunk_num, (chunk, vector) in enumerate(zip(chunks, vectors)):
            pid = int(hashlib.sha256(f"{path.name}:{chunk_num}".encode()).hexdigest()[:15], 16)
            points.append(PointStruct(
                id=pid,
                vector=np.asarray(vector).tolist(),
                payload={
                    "document_id": path.name,
                    "chunk_id": f"{path.name}#{chunk_num}",
                    "title": path.stem.replace("_", " "),
                    "department": department,
                    "allowed_groups": allowed_groups,
                    "sensitivity": sensitivity,
                    "source": f"local://{path.name}",
                    "text": chunk,
                    "version": 1,
                },
            ))
        if points:
            self.client.upsert(collection_name=COLLECTION, points=points)
        return len(points)

    def ingest_seed_documents(self, directory: str) -> int:
        path = Path(directory)
        total = 0
        for p in sorted(path.iterdir()):
            if p.is_file() and p.suffix.lower() in {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}:
                total += self.ingest_file(p)
        return total

    def retrieve(self, query: str, user_groups: list[str], top_k: int = 6) -> list[dict[str, Any]]:
        q = self.encoder.encode(f"query: {query}", normalize_embeddings=True).tolist()
        acl_filter = Filter(must=[FieldCondition(key="allowed_groups", match=MatchAny(any=user_groups))])
        result = self.client.query_points(
            collection_name=COLLECTION,
            query=q,
            query_filter=acl_filter,
            with_payload=True,
            limit=max(top_k * 5, 20),
        )
        candidates = []
        for dense_rank, point in enumerate(result.points, start=1):
            p = point.payload or {}
            text = str(p.get("text", ""))
            lexical = self._lexical_score(query, text)
            candidates.append({
                **p,
                "dense_score": float(point.score or 0.0),
                "lexical_score": lexical,
                "dense_rank": dense_rank,
            })
        # Reciprocal Rank Fusion over dense and lexical rankings.
        lexical_order = sorted(range(len(candidates)), key=lambda i: candidates[i]["lexical_score"], reverse=True)
        lexical_rank = {idx: rank for rank, idx in enumerate(lexical_order, start=1)}
        for idx, item in enumerate(candidates):
            item["rrf_score"] = 0.65 / (60 + item["dense_rank"]) + 0.35 / (60 + lexical_rank[idx])
        candidates.sort(key=lambda x: x["rrf_score"], reverse=True)
        return candidates[:top_k]

    @staticmethod
    def _lexical_score(query: str, text: str) -> float:
        q = set(re.findall(r"\w+", query.lower(), flags=re.UNICODE))
        t = set(re.findall(r"\w+", text.lower(), flags=re.UNICODE))
        return len(q & t) / max(1, len(q))
